import os
from openpyxl import Workbook, load_workbook
from datetime import datetime
from tqdm import tqdm

# Países correspondientes a cada zona
zonas_paises = {
    'Home': ['Argentina'],
    'Roaming': ['Paraguay', 'Estados Unidos', 'Canada', 'Panama', 'Brasil', 'Chile', 'Colombia', 'Republica Dominicana', 'El Salvador'],
    'Zona Util': ['Argentina', 'Paraguay', 'Estados Unidos', 'Canada', 'Panama', 'Brasil', 'Chile', 'Colombia', 'Republica Dominicana', 'El Salvador']
}

# Función para procesar el archivo seleccionado y generar el nuevo archivo
def procesar_archivo_zonas(ruta_archivo_seleccionado):
    # Cargar el archivo Excel original
    wb = load_workbook(ruta_archivo_seleccionado)
    sheet = wb.active

    # Crear un directorio para guardar el archivo de salida si no existe
    directorio_destino = r"C:\Migracion\Zone Management\Template de migracion"
    if not os.path.exists(directorio_destino):
        os.makedirs(directorio_destino)

    # Crear el nuevo archivo con el nombre correspondiente
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    archivo_destino = os.path.join(directorio_destino, f"Zone_management_{timestamp}.xlsx")

    # Crear un nuevo libro y hoja de trabajo
    wb_destino = Workbook()
    sheet_destino = wb_destino.active
    sheet_destino.title = "Migración"

    # Escribir la cabecera en el nuevo archivo
    sheet_destino['A1'] = "Account (M)"
    sheet_destino['B1'] = "Zone Name (M)"
    sheet_destino['C1'] = "Countries / Networks (M)"

    # Leer los valores desde la columna B (AccountName) desde la fila 2
    account_names = [sheet[f'B{fila}'].value for fila in range(2, sheet.max_row + 1) if sheet[f'B{fila}'].value]

    # Inicializar la fila destino y el contador total de filas
    fila_destino = 2
    total_filas = len(account_names) * sum(len(paises) for paises in zonas_paises.values())

    # Crear la barra de progreso
    with tqdm(total=total_filas, desc="Procesando cuentas", unit="fila") as pbar:
        for account_name in account_names:
            for zona, paises in zonas_paises.items():
                for pais in paises:
                    # Rellenar las columnas en el nuevo archivo
                    sheet_destino[f'A{fila_destino}'] = account_name
                    sheet_destino[f'B{fila_destino}'] = zona
                    sheet_destino[f'C{fila_destino}'] = pais

                    fila_destino += 1
                    pbar.update(1)

    # Guardar el nuevo archivo Excel generado
    wb_destino.save(archivo_destino)
    print(f"Archivo generado con éxito: {archivo_destino}")

# Main
if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Uso: python script_zonas.py <ruta_archivo_seleccionado>")
        sys.exit(1)

    ruta_archivo_seleccionado = sys.argv[1]

    if os.path.exists(ruta_archivo_seleccionado):
        procesar_archivo_zonas(ruta_archivo_seleccionado)
    else:
        print(f"El archivo especificado no existe: {ruta_archivo_seleccionado}")
