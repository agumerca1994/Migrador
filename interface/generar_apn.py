import os
import sys
from datetime import datetime
from openpyxl import load_workbook, Workbook

# Variables globales
archivo_generado_apn = ""  # Ruta del archivo generado como salida
selected_file = ""  # Archivo base seleccionado por la interfaz de usuario
archivo_bd = r"C:\Migracion\APN\bd.xlsx"  # Ruta del archivo BD con datos principales
output_directory = r"C:\Migracion\APN\Template de migracion"  # Directorio donde se guardará el archivo generado
cuentas_procesadas = []  # Lista para registrar las cuentas procesadas durante la ejecución


# Función para obtener la lista de "Accountnames" desde un archivo Excel
def obtener_accountnames(ruta_archivo):
    """
    Lee el archivo Excel especificado y obtiene los valores de la columna 'Accountname'.
    
    :param ruta_archivo: Ruta del archivo Excel que se analizará.
    :return: Lista de accountnames encontrados en el archivo.
    """
    wb = load_workbook(ruta_archivo, data_only=True)  # Carga el archivo Excel
    ws = wb.active  # Obtiene la hoja activa
    accountnames = []
    # Itera sobre la columna 2 desde la fila 2, obteniendo los valores de 'Accountname'
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        if row[0]:  # Si hay un valor no nulo, lo agrega a la lista
            accountnames.append(row[0])
    return accountnames


# Función para generar el archivo Excel de salida basado en coincidencias
def generar_apn_excel(coincidencias, archivo_bd, output_dir):
    """
    Genera un archivo Excel con los datos de las coincidencias entre el archivo base y la BD.
    
    :param coincidencias: Diccionario con las coincidencias encontradas.
    :param archivo_bd: Ruta del archivo BD con datos principales.
    :param output_dir: Directorio de salida para guardar el archivo generado.
    :return: Ruta del archivo generado.
    """
    global archivo_generado_apn, cuentas_procesadas

    # Verifica si el directorio de salida existe; si no, lo crea
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # Genera un nombre único para el archivo de salida basado en la fecha y hora actuales
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"APN_{timestamp}.xlsx")

    # Carga el archivo de BD y lee los datos
    bd_wb = load_workbook(archivo_bd, data_only=True)
    bd_ws = bd_wb.active

    # Crea un diccionario para almacenar los datos de la BD organizados por "Enterprise Name"
    bd_data = {
        row[1]: {
            "APN Type": row[10],  # Tipo de APN
            "Enterprise Name": row[1],  # Nombre de la empresa
            "APN #1": row[7],  # Primer APN
            "APN #2": row[8] if len(row) > 8 else None  # Segundo APN (opcional)
        }
        for row in bd_ws.iter_rows(min_row=2, values_only=True) if row[1]  # Ignora filas sin "Enterprise Name"
    }

    # Definición de la cabecera del archivo Excel generado
    cabecera = [
        "APN Type (M)", "Account (M)", "APN Id (M)", "APN Name (M)", "Description (M)", "EQOSID (O)",
        "Context ID (O)", "IP Address Type (M)", "Rating Group(M)", "HLR APN ID (O)", "MCC (O)", "MNC (O)",
        "HSS Profile Id (O)", "Profile Name 2G/3G (O)", "Bandwidth Uplink 2G/3G (O)",
        "Bandwidth Uplink Unit 2G/3G (O)", "Bandwidth Downlink 2G/3G (O)", "Bandwidth Downlink Unit 2G/3G (O)",
        "Profile Name 4G (O)", "Bandwidth Uplink 4G (O)", "Bandwidth Uplink Unit 4G (O)",
        "Bandwidth Downlink 4G(O)", "Bandwidth Downlink Unit 4G (O)"
    ]

    # Crea un nuevo archivo Excel y agrega la cabecera
    wb = Workbook()
    ws = wb.active
    ws.append(cabecera)

    apn_id = 1  # Contador para APN Id
    invalid_apn_values = {"", "---", None}  # Valores no válidos para APN

    # Procesa cada coincidencia y genera filas en el archivo Excel
    for i, account in enumerate(coincidencias.keys(), start=1):
        if account in bd_data:  # Si la cuenta está en los datos de la BD
            data = bd_data[account]
            cuentas_procesadas.append(account)  # Registra la cuenta como procesada
            # Muestra el progreso del procesamiento
            print(f"Procesando account_id {account} ({i}/{len(coincidencias)})...")
            # Agrega datos del primer APN si es válido
            if data["APN #1"] not in invalid_apn_values:
                ws.append([
                    data["APN Type"], data["Enterprise Name"], apn_id, data["APN #1"], "", "", "", "IPv4", "AQ-121", "", "", "", "", "", "", "", "Kbps", "", "", "Kbps"
                ])
                apn_id += 1
            # Agrega datos del segundo APN si es válido
            if data["APN #2"] not in invalid_apn_values:
                ws.append([
                    data["APN Type"], data["Enterprise Name"], apn_id, data["APN #2"], "", "", "", "IPv4", "AQ-121", "", "", "", "", "", "", "", "Kbps", "", "", "Kbps"
                ])
                apn_id += 1

    # Guarda el archivo generado
    wb.save(output_path)
    archivo_generado_apn = output_path  # Actualiza la variable global con la ruta del archivo generado
    return output_path


# Función principal para procesar el archivo seleccionado
def procesar_archivo():
    """
    Procesa el archivo seleccionado, encuentra coincidencias y genera un archivo de salida.
    """
    global selected_file, archivo_bd, output_directory

    if not selected_file:  # Verifica si se ha seleccionado un archivo
        print("No se ha seleccionado un archivo base.")
        return

    if not os.path.exists(archivo_bd):  # Verifica si el archivo BD existe
        print("El archivo bd.xlsx no se encuentra en el directorio especificado.")
        return

    # Obtiene los accountnames del archivo base y del archivo BD
    accountnames_base = obtener_accountnames(selected_file)
    accountnames_bd = obtener_accountnames(archivo_bd)

    # Encuentra coincidencias entre el archivo base y el archivo BD
    coincidencias = {}
    for account in accountnames_base:
        if account in accountnames_bd:
            coincidencias[account] = [account]

    # Genera el archivo de salida si hay coincidencias
    if coincidencias:
        print(f"Iniciando la consulta de {len(coincidencias)} account_ids...")
        output_file = generar_apn_excel(coincidencias, archivo_bd, output_directory)
        print(f"Los resultados se han guardado en '{output_file}'.")
    else:
        print("No se encontraron coincidencias entre los Accountname.")


# Punto de entrada principal
if __name__ == "__main__":
    if len(sys.argv) > 1:  # Verifica si se pasó un archivo como argumento
        selected_file = sys.argv[1]
    else:
        print("Error: No se proporcionó un archivo base. Usa: python script.py <ruta_del_archivo>")
        sys.exit(1)

    procesar_archivo()
